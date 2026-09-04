from docx import Document
from docx.shared import Pt, Inches
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re


# ============================================================
# CONFIG
# ============================================================

SOURCE_MD = "/mnt/data/Pasted markdown.md"
OUTPUT_DOCX = "/mnt/data/converted_google_docs_style.docx"


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_markdown(text):
    """
    Remove Markdown presentation syntax while preserving
    the actual content.
    """

    text = text.replace("\\|", "|")
    text = text.replace("\\", "")
    text = text.replace("**", "")
    text = text.replace("`", "")
    text = text.replace("\u00a0", " ")

    return text.strip()


# ============================================================
# TABLE PARSING
# ============================================================

def parse_table_row(line):
    """
    Convert a Markdown table row into a list of cells.

    Handles escaped pipes such as:
        \\| Keyword \\| Intent \\|
    """

    line = line.strip()

    if line.startswith("\\|"):
        line = line[1:]

    line = line.replace("\\|", "|")

    return [
        cell.strip()
        for cell in line.strip("|").split("|")
    ]


def is_table_separator(line):
    """
    Detect Markdown separator rows such as:

    |---|---|---|
    |:---|---:|---|
    """

    cells = parse_table_row(line)

    if not cells:
        return False

    return all(
        re.fullmatch(r":?-+:?", cell.replace(" ", ""))
        for cell in cells
    )


# ============================================================
# WORD TABLE FORMATTING
# ============================================================

def shade_cell(cell, color="E7E6E6"):
    tc_pr = cell._tc.get_or_add_tcPr()

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)

    tc_pr.append(shading)


def set_cell_margins(
    cell,
    top=80,
    start=90,
    bottom=80,
    end=90
):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ]:

        node = tc_mar.find(qn(f"w:{margin}"))

        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)

        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    """
    Clean, subtle borders similar to a normal Google Docs
    assignment table.
    """

    tbl_pr = table._tbl.tblPr

    borders = tbl_pr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ]:

        element = OxmlElement(f"w:{edge}")

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")

        borders.append(element)


def repeat_table_header(row):
    """
    Makes the first table row repeat when a table spans pages.
    """

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")

    tr_pr.append(tbl_header)


def write_cell(
    cell,
    text,
    bold=False,
    font_size=8.5
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(
        clean_markdown(text)
    )

    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.bold = bold

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    set_cell_margins(cell)


# ============================================================
# ADD WORD TABLE
# ============================================================

def add_word_table(document, rows):

    if not rows:
        return

    column_count = len(rows[0])

    table = document.add_table(
        rows=1,
        cols=column_count
    )

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    set_table_borders(table)

    # ----------------------------
    # Header
    # ----------------------------

    header = table.rows[0]

    repeat_table_header(header)

    for column, value in enumerate(rows[0]):

        cell = header.cells[column]

        write_cell(
            cell,
            value,
            bold=True,
            font_size=8.5
        )

        shade_cell(cell)

    # ----------------------------
    # Body
    # ----------------------------

    for row_data in rows[1:]:

        cells = table.add_row().cells

        row_data = (
            row_data
            + [""] * column_count
        )[:column_count]

        for column, value in enumerate(row_data):

            write_cell(
                cells[column],
                value,
                bold=False,
                font_size=8.5
            )

    # ----------------------------
    # Row behavior
    # ----------------------------

    for row in table.rows:

        row.height_rule = (
            WD_ROW_HEIGHT_RULE.AT_LEAST
        )

    document.add_paragraph()


# ============================================================
# DOCUMENT STYLING
# ============================================================

def configure_styles(document):

    styles = document.styles

    for style_name in [
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]:

        styles[style_name].font.name = "Arial"

    # Body
    styles["Normal"].font.size = Pt(10.5)

    # Document title
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True

    # Major sections
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True

    # Subsections
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True

    # Smaller headings
    styles["Heading 3"].font.size = Pt(11)
    styles["Heading 3"].font.bold = True

    # Spacing
    for style_name in [
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]:

        style = styles[style_name]

        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(6)


# ============================================================
# MAIN CONVERSION
# ============================================================

def convert_markdown_to_docx():

    # ----------------------------
    # Read Markdown
    # ----------------------------

    with open(
        SOURCE_MD,
        "r",
        encoding="utf-8"
    ) as file:

        markdown = file.read()

    lines = markdown.splitlines()

    # ----------------------------
    # Create document
    # ----------------------------

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    configure_styles(document)

    # ----------------------------
    # Process lines
    # ----------------------------

    i = 0

    while i < len(lines):

        line = lines[i].strip()

        # Empty line
        if not line:

            i += 1
            continue

        # ====================================================
        # TABLE DETECTION
        # ====================================================

        if line.startswith("\\|"):

            # Look ahead for separator row
            j = i + 1

            while (
                j < len(lines)
                and not lines[j].strip()
            ):
                j += 1

            if (
                j < len(lines)
                and lines[j].strip().startswith("\\|")
                and is_table_separator(lines[j])
            ):

                rows = []

                # Header
                rows.append(
                    parse_table_row(line)
                )

                # Skip separator
                i = j + 1

                # Read table rows.
                #
                # IMPORTANT:
                # The original Markdown may contain blank
                # lines between rows, so we deliberately
                # skip blank lines here.

                while i < len(lines):

                    k = i

                    while (
                        k < len(lines)
                        and not lines[k].strip()
                    ):
                        k += 1

                    if (
                        k < len(lines)
                        and lines[k].strip().startswith("\\|")
                    ):

                        rows.append(
                            parse_table_row(lines[k])
                        )

                        i = k + 1

                    else:

                        i = k
                        break

                add_word_table(
                    document,
                    rows
                )

                continue

        # ====================================================
        # HEADINGS
        # ====================================================

        if line.startswith("**# "):

            paragraph = document.add_paragraph(
                style="Title"
            )

            paragraph.add_run(
                clean_markdown(line[2:])
            )

            paragraph.paragraph_format.space_after = Pt(12)

        elif line.startswith("**## "):

            paragraph = document.add_paragraph(
                style="Heading 1"
            )

            paragraph.add_run(
                clean_markdown(line[3:])
            )

            paragraph.paragraph_format.keep_with_next = True

        elif line.startswith("**### "):

            paragraph = document.add_paragraph(
                style="Heading 2"
            )

            paragraph.add_run(
                clean_markdown(line[4:])
            )

            paragraph.paragraph_format.keep_with_next = True

        elif line.startswith("**#### "):

            paragraph = document.add_paragraph(
                style="Heading 3"
            )

            paragraph.add_run(
                clean_markdown(line[5:])
            )

        # ====================================================
        # BULLET LISTS
        # ====================================================

        elif line.startswith("\\- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                clean_markdown(line[2:])
            )

        elif line.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                clean_markdown(line[2:])
            )

        # ====================================================
        # NUMBERED LISTS
        # ====================================================

        elif re.match(
            r"^\d+\.\s",
            line
        ):

            paragraph = document.add_paragraph(
                style="List Number"
            )

            content = re.sub(
                r"^\d+\.\s",
                "",
                line
            )

            paragraph.add_run(
                clean_markdown(content)
            )

        # ====================================================
        # HORIZONTAL RULE
        # ====================================================

        elif re.fullmatch(
            r"[-*]{3,}",
            line
        ):

            document.add_paragraph()

        # ====================================================
        # NORMAL PARAGRAPH
        # ====================================================

        else:

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(5)

            paragraph.add_run(
                clean_markdown(line)
            )

        i += 1

    # ----------------------------
    # Save
    # ----------------------------

    document.save(
        OUTPUT_DOCX
    )

    # ----------------------------
    # Verification
    # ----------------------------

    verification = Document(
        OUTPUT_DOCX
    )

    print(
        "Created:",
        OUTPUT_DOCX
    )

    print(
        "Number of real Word tables:",
        len(verification.tables)
    )

    for number, table in enumerate(
        verification.tables,
        start=1
    ):

        print(
            f"Table {number}: "
            f"{len(table.rows)} rows x "
            f"{len(table.columns)} columns"
        )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    convert_markdown_to_docx()